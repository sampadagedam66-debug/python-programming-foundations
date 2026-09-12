"""
excel_write.py

Topic: Writing Excel Files in Python (using openpyxl)
Covers: creating a workbook & writing cells | writing rows with .append() |
        basic cell formatting (bold headers)
Mini use-case: an employee salary sheet generator
Requires: pip install openpyxl
"""

from openpyxl import Workbook
from openpyxl.styles import Font

# ---------------------------------------------------------
# SECTION 1: Creating a Workbook & Writing Cells — Monthly Budget
# ---------------------------------------------------------
print("---- Writing Cells: Monthly Budget Sheet ----")

budget_workbook = Workbook()
budget_sheet = budget_workbook.active
budget_sheet.title = "Budget"

# writing individual cells using (row, column) references
budget_sheet.cell(row=1, column=1, value="Category")
budget_sheet.cell(row=1, column=2, value="Amount")
budget_sheet.cell(row=2, column=1, value="Rent")
budget_sheet.cell(row=2, column=2, value=12000)
budget_sheet.cell(row=3, column=1, value="Groceries")
budget_sheet.cell(row=3, column=2, value=4500)

budget_workbook.save("monthly_budget.xlsx")
print("Budget sheet saved to monthly_budget.xlsx")


# ---------------------------------------------------------
# SECTION 2: Writing Rows With .append() — Tournament Results
# ---------------------------------------------------------
print("\n---- .append(): Tournament Results ----")

tournament_workbook = Workbook()
tournament_sheet = tournament_workbook.active
tournament_sheet.title = "Results"

tournament_sheet.append(["Team", "Matches Won", "Matches Lost"])
tournament_sheet.append(["Titans", 8, 2])
tournament_sheet.append(["Warriors", 6, 4])
tournament_sheet.append(["Strikers", 9, 1])

tournament_workbook.save("tournament_results.xlsx")
print("Tournament results saved to tournament_results.xlsx")


# ---------------------------------------------------------
# SECTION 3: Mini Project — Employee Salary Sheet Generator
# ---------------------------------------------------------
# Combines cell writing, .append(), and basic formatting —
# a different domain from both sections above.

print("\n---- Mini Project: Employee Salary Sheet ----")

employees = [
    ("Ritika Verma", "Software Engineer", 55000),
    ("Devraj Singh", "QA Engineer", 42000),
    ("Meera Nair", "UI/UX Designer", 48000),
]

salary_workbook = Workbook()
salary_sheet = salary_workbook.active
salary_sheet.title = "Salaries"

headers = ["Employee Name", "Role", "Monthly Salary"]
salary_sheet.append(headers)

# make the header row bold — a small but noticeable formatting touch
for column_number in range(1, len(headers) + 1):
    header_cell = salary_sheet.cell(row=1, column=column_number)
    header_cell.font = Font(bold=True)

for employee_name, role, salary in employees:
    salary_sheet.append([employee_name, role, salary])

salary_workbook.save("employee_salaries.xlsx")
print("Employee salary sheet saved to employee_salaries.xlsx (with bold headers)")
