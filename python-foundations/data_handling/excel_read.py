"""
excel_read.py

Topic: Reading Excel Files in Python (using openpyxl)
Covers: reading cells by reference | iterating rows with .iter_rows() |
        computing a total from read data
Mini use-case: computing total payroll from the salary sheet
Requires: pip install openpyxl
Note: run excel_write.py first, since this file reads the .xlsx files it creates.
"""

from openpyxl import load_workbook

# ---------------------------------------------------------
# SECTION 1: Reading Cells by Reference — Monthly Budget
# ---------------------------------------------------------
print("---- Reading Cells: Monthly Budget Sheet ----")

budget_workbook = load_workbook("monthly_budget.xlsx")
budget_sheet = budget_workbook.active

category_1 = budget_sheet.cell(row=2, column=1).value
amount_1 = budget_sheet.cell(row=2, column=2).value
category_2 = budget_sheet.cell(row=3, column=1).value
amount_2 = budget_sheet.cell(row=3, column=2).value

print(f"{category_1}: Rs {amount_1}")
print(f"{category_2}: Rs {amount_2}")


# ---------------------------------------------------------
# SECTION 2: Iterating Rows With .iter_rows() — Tournament Results
# ---------------------------------------------------------
print("\n---- .iter_rows(): Tournament Results ----")

tournament_workbook = load_workbook("tournament_results.xlsx")
tournament_sheet = tournament_workbook.active

most_wins = 0
leading_team = ""

for row in tournament_sheet.iter_rows(min_row=2, values_only=True):
    team_name, matches_won, matches_lost = row
    print(f"{team_name}: {matches_won} won, {matches_lost} lost")
    if matches_won > most_wins:
        most_wins = matches_won
        leading_team = team_name

print(f"\nTeam with most wins: {leading_team} ({most_wins} wins)")


# ---------------------------------------------------------
# SECTION 3: Mini Project — Total Payroll Calculator
# ---------------------------------------------------------
# Reads the salary sheet created by excel_write.py and computes
# the company's total monthly payroll — a different domain from
# both sections above.

print("\n---- Mini Project: Total Payroll Calculator ----")

salary_workbook = load_workbook("employee_salaries.xlsx")
salary_sheet = salary_workbook.active

total_payroll = 0

for row in salary_sheet.iter_rows(min_row=2, values_only=True):
    employee_name, role, salary = row
    print(f"{employee_name} ({role}): Rs {salary}")
    total_payroll += salary

print(f"\nTotal monthly payroll: Rs {total_payroll}")
